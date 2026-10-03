"""The export keeps persisted evidence visible and untrusted text inert."""

from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.models.research import ResearchRun, ResearchStatus
from app.models.research_review import ResearchReportReview, ResearchReviewDecision
from app.schemas.research import ResearchEvidence, ResearchReport
from app.services.research_agent import validate_report_sources
from app.services.research_export import export_research_report


def sample_run(
    *,
    title: str = "商品标题",
    summary: str = "报告摘要",
    claim: str = "价格有优势",
    original_name: str = "资料.txt",
    excerpt: str = "已验证摘录",
) -> tuple[ResearchRun, ResearchReportReview]:
    source_id = "kb:2:document:3:chunk:0"
    run = ResearchRun(
        id=7,
        product_id=1,
        requested_by_id=4,
        goal="竞品比较",
        knowledge_base_ids=[2],
        product_snapshot={"sku": "SKU-1", "title": title},
        status=ResearchStatus.SUCCESS,
        provider="test",
        model="research-model",
        prompt_version="research-v1",
        steps=[],
        evidence=[
            {
                "source_id": source_id,
                "source_type": "knowledge",
                "text": excerpt,
                "knowledge_base_id": 2,
                "document_id": 3,
                "chunk_id": "document:3:chunk:0",
                "chunk_index": 0,
                "page_number": 2,
                "original_name": original_name,
                "distance": 0.2,
            }
        ],
        report={
            "outcome": "supported",
            "summary": summary,
            "findings": [{"claim": claim, "source_ids": [source_id]}],
            "recommendations": [
                {"action": "保留定价", "reason": "竞争力", "source_ids": [source_id]}
            ],
            "evidence_gaps": [],
        },
    )
    run.evidence.append(
        {
            "source_id": "product:1:snapshot",
            "source_type": "product",
            "product_id": 1,
            "text": "商品快照",
        }
    )
    run.report = validate_report_sources(
        ResearchReport.model_validate(run.report),
        [ResearchEvidence.model_validate(item) for item in run.evidence],
    )
    review = ResearchReportReview(
        id=9,
        run_id=7,
        reviewed_by_id=5,
        decision=ResearchReviewDecision.APPROVED,
        comment=None,
    )
    return run, review


def test_export_contains_report_and_registered_sources() -> None:
    run, review = sample_run()
    workbook = load_workbook(BytesIO(export_research_report(run, review)))
    assert workbook.sheetnames == ["概览", "发现与建议", "证据"]
    all_values = [cell.value for sheet in workbook for row in sheet for cell in row]
    for expected in (
        "SKU-1",
        "商品标题",
        "报告摘要",
        "价格有优势",
        "保留定价",
        "kb:2:document:3:chunk:0",
        "资料.txt",
        "已验证摘录",
        "结论基于已上传且经程序核验的资料，仍需业务复核",
    ):
        assert expected in all_values
    assert 7 in all_values
    assert 5 in all_values


def test_export_never_writes_untrusted_formula() -> None:
    run, review = sample_run(
        title="=1+1",
        summary=" +cmd",
        claim="-2",
        original_name="@SUM(1)",
        excerpt='  =HYPERLINK("https://example.invalid")',
    )
    workbook = load_workbook(BytesIO(export_research_report(run, review)))
    untrusted_cells = [
        cell
        for sheet in workbook
        for row in sheet
        for cell in row
        if isinstance(cell.value, str)
        and any(marker in cell.value for marker in ("=1+1", "+cmd", "-2", "@SUM(1)", "=HYPERLINK"))
    ]
    assert len(untrusted_cells) == 5
    assert all(cell.data_type != "f" and cell.hyperlink is None for cell in untrusted_cells)


def test_export_handles_xlsx_illegal_control_characters() -> None:
    run, review = sample_run(title="商品\x00标题", excerpt="来源\x01摘录")
    workbook = load_workbook(BytesIO(export_research_report(run, review)))
    values = [cell.value for sheet in workbook for row in sheet for cell in row]
    assert "商品标题" in values
    assert "来源摘录" in values
