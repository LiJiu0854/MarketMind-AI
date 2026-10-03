"""从已批准研究的持久化快照生成内存 Excel 交付件。"""

from io import BytesIO

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]
from pydantic import ValidationError

from app.core.errors import AppError
from app.models.research import ResearchRun
from app.models.research_review import ResearchReportReview
from app.services.research_review import validate_persisted_research

DISCLAIMER = "结论基于已上传且经程序核验的资料，仍需业务复核"


def _safe_excel_text(value: str) -> str:
    """Dangerous prefixes stay text even after leading spaces."""
    value = "".join(
        char
        for char in value
        if char in "\t\n\r"
        or 0x20 <= ord(char) <= 0xD7FF
        or 0xE000 <= ord(char) <= 0xFFFD
        or 0x10000 <= ord(char) <= 0x10FFFF
    )
    return f"'{value}" if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _safe(value: object) -> object:
    return _safe_excel_text(value) if isinstance(value, str) else value


def export_research_report(run: ResearchRun, review: ResearchReportReview) -> bytes:
    try:
        report, evidence = validate_persisted_research(run)
    except (ValidationError, ValueError):
        raise AppError("RESEARCH_EXPORT_INVALID", "研究报告或来源数据无效", 409) from None

    workbook = Workbook()
    overview = workbook.active
    assert overview is not None
    overview.title = "概览"
    overview.append(["字段", "内容"])
    overview_rows: list[tuple[str, object]] = [
        ("研究 ID", run.id),
        ("商品 SKU", run.product_snapshot.get("sku")),
        ("商品标题", run.product_snapshot.get("title")),
        ("研究目标", run.goal),
        ("报告结论", report.outcome),
        ("报告摘要", report.summary),
        ("审核决定", review.decision.value),
        ("审核人 ID", review.reviewed_by_id),
        ("审核时间", review.reviewed_at),
        ("使用说明", DISCLAIMER),
    ]
    for label, value in overview_rows:
        overview.append([label, _safe(value)])

    findings = workbook.create_sheet("发现与建议")
    findings.append(["类型", "内容", "理由", "来源 ID"])
    for finding in report.findings:
        findings.append(
            ["发现", _safe(finding.claim), None, _safe_excel_text(", ".join(finding.source_ids))]
        )
    for recommendation in report.recommendations:
        findings.append(
            [
                "建议",
                _safe(recommendation.action),
                _safe(recommendation.reason),
                _safe_excel_text(", ".join(recommendation.source_ids)),
            ]
        )

    sources = workbook.create_sheet("证据")
    sources.append(["来源 ID", "类型", "文件名", "页码", "片段序号", "片段 ID", "摘录"])
    for source in evidence:
        sources.append(
            [
                _safe(source.source_id),
                source.source_type,
                _safe(source.original_name),
                source.page_number,
                source.chunk_index,
                _safe(source.chunk_id),
                _safe(source.text),
            ]
        )

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column, width in {
            "A": 25,
            "B": 55,
            "C": 35,
            "D": 20,
            "E": 20,
            "F": 35,
            "G": 80,
        }.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
